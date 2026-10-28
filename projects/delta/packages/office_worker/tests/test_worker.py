"""The office worker speaks the capability ABI and stages only what it has verified."""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest

from packages.office_worker import worker

WORKER = Path(worker.__file__).resolve()


@pytest.fixture
def staging(tmp_path: Path) -> Path:
    directory = tmp_path / "staging"
    directory.mkdir()
    return directory


def _job(
    staging: Path | None, capability: str = "office.docx.create", **overrides: object
) -> dict:
    job: dict[str, object] = {
        "abi_version": 2,
        "capability_id": capability,
        "job_id": "job-office-1",
        "input_files": [],
        "arguments": {
            "path": "reports/q3.docx",
            "content": {
                "title": "Q3",
                "blocks": [{"type": "paragraph", "text": "Done."}],
            },
        },
        "artifact_staging_dir": str(staging) if staging is not None else None,
    }
    job.update(overrides)
    return job


def test_a_docx_job_stages_a_verified_candidate(staging: Path) -> None:
    result = worker.run_job(_job(staging))

    assert result["state"] == "completed"
    artifact = result["artifacts"][0]
    assert artifact["staging_path"] == str(staging / "candidate.docx")
    assert artifact["relative_path"] == "reports/q3.docx"
    assert artifact["kind"] == "file"
    assert artifact["size"] == (staging / "candidate.docx").stat().st_size
    assert result["result"]["format"] == "docx"
    assert result["result"]["verification"]["status"] == "passed"
    assert result["result"]["media_type"].endswith("wordprocessingml.document")


def test_the_result_names_the_spec_it_was_built_from(staging: Path) -> None:
    first = worker.run_job(_job(staging))["result"]["spec_sha256"]
    changed = _job(staging)
    changed["arguments"]["content"]["blocks"][0]["text"] = "Different."

    assert worker.run_job(changed)["result"]["spec_sha256"] != first


@pytest.mark.parametrize(
    ("capability", "content", "kind"),
    [
        (
            "office.pptx.create",
            {
                "title": "Deck",
                "slides": [{"layout": "bullets", "title": "A", "items": ["x"]}],
            },
            "pptx",
        ),
        (
            "office.xlsx.create",
            {
                "title": "Book",
                "sheets": [{"name": "S", "columns": ["x"], "rows": [[1]]}],
            },
            "xlsx",
        ),
    ],
)
def test_each_format_has_a_capability(
    staging: Path, capability: str, content: dict, kind: str
) -> None:
    job = _job(staging, capability)
    job["arguments"] = {"path": f"out/file.{kind}", "content": content}

    result = worker.run_job(job)

    assert result["state"] == "completed"
    assert result["result"]["format"] == kind


def test_a_refused_spec_is_failed_with_a_code_and_nothing_is_staged_as_an_artifact(
    staging: Path,
) -> None:
    job = _job(staging)
    job["arguments"]["content"] = {
        "title": "t",
        "blocks": [{"type": "heading", "level": 9, "text": "x"}],
    }

    result = worker.run_job(job)

    assert result["state"] == "failed"
    assert result["artifacts"] == []
    assert result["diagnostics"]["error_code"] == "invalid_spec"
    assert "heading level" in result["diagnostics"]["error_message"]


def test_missing_staging_is_refused(staging: Path) -> None:
    result = worker.run_job(_job(None))

    assert result["diagnostics"]["error_code"] == "staging"


def test_staging_that_is_not_a_directory_is_refused(tmp_path: Path) -> None:
    result = worker.run_job(_job(tmp_path / "missing"))

    assert result["diagnostics"]["error_code"] == "staging"


def test_a_relative_staging_path_is_refused(staging: Path) -> None:
    result = worker.run_job(_job(Path("relative-staging")))

    assert result["diagnostics"]["error_code"] == "staging"


def test_a_job_without_a_path_is_refused(staging: Path) -> None:
    job = _job(staging)
    job["arguments"] = {"content": {"title": "t", "blocks": []}}

    assert worker.run_job(job)["diagnostics"]["error_code"] == "arguments"


def test_a_job_with_a_wrong_abi_version_is_refused(staging: Path) -> None:
    assert (
        worker.run_job(_job(staging, abi_version=1))["diagnostics"]["error_code"]
        == "abi_mismatch"
    )


def test_an_unknown_capability_is_refused(staging: Path) -> None:
    assert (
        worker.run_job(_job(staging, "office.pdf.create"))["diagnostics"]["error_code"]
        == "unknown_capability"
    )


def test_a_candidate_that_does_not_match_its_report_is_refused(
    staging: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def lying_builder(content: dict, path: Path) -> dict:
        path.write_bytes(b"")  # a file that is not a document at all
        return {"tables": 0}

    monkeypatch.setitem(worker.BUILDERS, "docx", lying_builder)

    result = worker.run_job(_job(staging))

    assert result["state"] == "failed"
    assert result["diagnostics"]["error_code"] == "verify_failed"
    assert result["artifacts"] == []


def test_the_protocol_round_trips_through_a_subprocess(staging: Path) -> None:
    job = _job(staging)
    stdin = json.dumps(job) + "\n" + "\n"  # the job, then the empty secret line
    completed = subprocess.run(
        [sys.executable, "-I", str(WORKER)],
        input=stdin,
        capture_output=True,
        text=True,
        timeout=120,
        check=True,
    )

    result = json.loads(completed.stdout.strip().splitlines()[-1])
    assert result["job_id"] == "job-office-1"
    assert result["state"] == "completed"
    assert (staging / "candidate.docx").is_file()


def test_a_path_with_the_wrong_extension_is_refused(staging: Path) -> None:
    job = _job(staging, "office.xlsx.create")
    job["arguments"] = {
        "path": "reports/q3.docx",
        "content": {"title": "t", "sheets": []},
    }

    result = worker.run_job(job)

    assert result["diagnostics"]["error_code"] == "arguments"
    assert "must end with .xlsx" in result["diagnostics"]["error_message"]


def _spec_with_placeholders() -> dict:
    return {
        "title": "Invoice {{number}}",
        "blocks": [
            {"type": "paragraph", "text": "Customer: {{customer}}"},
            {
                "type": "table",
                "columns": ["Item", "Amount"],
                "rows": [["Support", "{{amount}}"]],
            },
        ],
    }


def test_placeholders_are_filled_from_data(staging: Path) -> None:
    job = _job(staging)
    job["arguments"]["content"] = _spec_with_placeholders()
    job["arguments"]["data"] = {"number": "1042", "customer": "Acme", "amount": 1200}

    result = worker.run_job(job)

    assert result["state"] == "completed", result["diagnostics"]
    assert result["result"]["counts"]["tables"] == 1
    from docx import Document

    document = Document(str(staging / "candidate.docx"))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Invoice 1042" in text
    assert "Customer: Acme" in text
    assert "{{" not in text
    assert document.tables[0].rows[1].cells[1].text == "1200"


def test_a_placeholder_without_a_value_is_refused_and_leaves_no_file(staging: Path) -> None:
    job = _job(staging)
    job["arguments"]["content"] = _spec_with_placeholders()
    job["arguments"]["data"] = {"number": "1042", "customer": "Acme"}

    result = worker.run_job(job)

    assert result["state"] == "failed"
    assert result["diagnostics"]["error_code"] == "missing_value"
    assert '"amount"' in result["diagnostics"]["error_message"]
    assert not (staging / "candidate.docx").exists()


def test_a_placeholder_without_any_data_is_refused(staging: Path) -> None:
    job = _job(staging)
    job["arguments"]["content"] = _spec_with_placeholders()

    result = worker.run_job(job)

    assert result["state"] == "failed"
    assert result["diagnostics"]["error_code"] == "missing_value"


def test_a_dry_run_checks_the_spec_and_registers_nothing(staging: Path) -> None:
    job = _job(staging)
    job["arguments"]["dry_run"] = True

    result = worker.run_job(job)

    assert result["state"] == "completed"
    assert result["artifacts"] == []
    assert result["result"]["dry_run"] is True
    assert result["result"]["verification"]["status"] == "passed"
    assert not (staging / "candidate.docx").exists()


def test_a_dry_run_reports_an_invalid_spec_with_the_valid_range(staging: Path) -> None:
    job = _job(staging)
    job["arguments"]["dry_run"] = True
    job["arguments"]["content"] = {"title": "Q3", "blocks": [{"type": "table", "columns": [], "rows": []}]}

    result = worker.run_job(job)

    assert result["state"] == "failed"
    message = result["diagnostics"]["error_message"]
    assert "needs at least one column" in message
    assert "1 to 64" in message


def test_the_data_argument_must_be_an_object(staging: Path) -> None:
    job = _job(staging)
    job["arguments"]["data"] = ["not", "an", "object"]

    result = worker.run_job(job)

    assert result["state"] == "failed"
    assert result["diagnostics"]["error_code"] == "arguments"
