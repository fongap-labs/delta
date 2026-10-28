"""Native charts in XLSX and PPTX: built from the spec, and refused when the data cannot chart."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from packages.office_worker import pptx_build, spec, xlsx_build


def _sales_sheet(**overrides: object) -> dict:
    sheet: dict[str, object] = {
        "name": "Sales",
        "columns": ["Region", "Q1", "Q2"],
        "rows": [["North", 120, 140], ["South", 90, 110], ["East", 75, 80]],
        "charts": [
            {
                "type": "column",
                "title": "Sales by region",
                "category_column": "Region",
                "value_columns": ["Q1", "Q2"],
            }
        ],
    }
    sheet.update(overrides)
    return sheet


def test_an_xlsx_chart_is_a_native_chart_part(tmp_path: Path) -> None:
    path = tmp_path / "sales.xlsx"
    counts = xlsx_build.build(
        {"title": "Sales", "sheets": [_sales_sheet(), _sales_sheet(name="Trend", charts=[{"type": "line", "category_column": "Region", "value_columns": ["Q2"]}])]},
        path,
    )

    assert counts["charts"] == 2
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    assert "xl/charts/chart1.xml" in names
    assert "xl/charts/chart2.xml" in names


def test_a_pie_chart_takes_one_value_column(tmp_path: Path) -> None:
    sheet = _sales_sheet(charts=[{"type": "pie", "category_column": "Region", "value_columns": ["Q1", "Q2"]}])

    with pytest.raises(ValueError, match="pie charts take exactly one value column"):
        xlsx_build.build({"title": "Sales", "sheets": [sheet]}, tmp_path / "bad.xlsx")


def test_a_chart_over_a_text_column_is_refused(tmp_path: Path) -> None:
    sheet = _sales_sheet(charts=[{"type": "column", "category_column": "Q1", "value_columns": ["Region"]}])

    with pytest.raises(ValueError, match="must be a number or a formula"):
        xlsx_build.build({"title": "Sales", "sheets": [sheet]}, tmp_path / "bad.xlsx")


def test_a_chart_needs_a_column_the_sheet_has(tmp_path: Path) -> None:
    sheet = _sales_sheet(charts=[{"type": "column", "category_column": "Region", "value_columns": ["Missing"]}])

    with pytest.raises(ValueError, match="must be a different column of the sheet"):
        xlsx_build.build({"title": "Sales", "sheets": [sheet]}, tmp_path / "bad.xlsx")


def test_a_pptx_chart_slide_holds_the_series_as_data(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    counts = pptx_build.build(
        {
            "title": "Deck",
            "slides": [
                {
                    "layout": "chart",
                    "title": "Revenue",
                    "chart": {
                        "type": "line",
                        "title": "By quarter",
                        "categories": ["Q1", "Q2", "Q3"],
                        "series": [{"name": "2025", "values": [1.5, 2, 2.75]}],
                    },
                }
            ],
        },
        path,
    )

    assert counts["charts"] == 1
    slide = Presentation(str(path)).slides[0]
    chart_shapes = [shape for shape in slide.shapes if shape.has_chart]
    assert len(chart_shapes) == 1
    chart = chart_shapes[0].chart
    assert list(chart.plots[0].categories) == ["Q1", "Q2", "Q3"]
    assert list(chart.series[0].values) == [1.5, 2.0, 2.75]


def test_a_chart_with_mismatched_values_says_what_to_change() -> None:
    with pytest.raises(ValueError, match="give one value per category"):
        spec.chart(
            {"type": "column", "categories": ["A", "B"], "series": [{"name": "x", "values": [1]}]},
            "chart",
        )


def test_a_non_finite_chart_value_is_refused() -> None:
    with pytest.raises(ValueError, match="must be a finite number"):
        spec.chart(
            {"type": "column", "categories": ["A"], "series": [{"name": "x", "values": [float("nan")]}]},
            "chart",
        )
