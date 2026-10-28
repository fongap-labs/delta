"""Builds a finished DOCX from a structured spec (ADR-0085).

The house style is set once here, on the document's own styles, so every heading, paragraph, list
and table in the file shares one typography and one palette. Tables repeat their header row on each
page, numbers are aligned to the right, and the footer carries a live page number.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from packages.office_worker import chart_render, spec

BODY_FONT = "Calibri"
HEADING_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
TABLE_STYLE = "Light Grid Accent 1"
HEADING_LEVELS = (1, 2, 3)


def is_heading(paragraph: Any) -> bool:
    """True for a paragraph in one of the heading styles this builder writes."""
    style = paragraph.style
    return style is not None and str(style.name or "").startswith("Heading")


def _style_document(document: Any) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = BODY_FONT
        style.font.color.rgb = HEADING_COLOR
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.5)


def _page_number_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    for tag, text in (
        ("begin", None),
        ("instr", " PAGE "),
        ("separate", None),
        ("end", None),
    ):
        if tag == "instr":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        run._r.append(element)


def _header_and_footer(document: Any, header: str, footer: str) -> None:
    for section in document.sections:
        if header:
            section.header.paragraphs[0].text = header
        footer_paragraph = section.footer.paragraphs[0]
        footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer:
            footer_paragraph.add_run(f"{footer}    ")
        _page_number_field(footer_paragraph)


def _repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _table(document: Any, block: dict[str, Any], label: str) -> None:
    headers, rows = spec.table(block.get("columns"), block.get("rows"), label)
    if block.get("caption"):
        caption = document.add_paragraph()
        caption.add_run(
            spec.text(block["caption"], f"{label} caption", max_chars=512)
        ).bold = True
    grid = document.add_table(rows=1, cols=len(headers))
    grid.style = TABLE_STYLE
    grid.autofit = True
    for index, name in enumerate(headers):
        cell = grid.rows[0].cells[index]
        cell.text = name
        cell.paragraphs[0].runs[0].bold = True
    _repeat_header(grid.rows[0])
    for row in rows:
        cells = grid.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = spec.display(value)
            if spec.is_number(value):
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph()


def _block(document: Any, block: dict[str, Any], index: int) -> dict[str, int]:
    label = f"blocks[{index}]"
    kind = block.get("type")
    if kind == "heading":
        level = block.get("level", 1)
        if level not in HEADING_LEVELS:
            raise ValueError(f"{label} heading level must be 1, 2 or 3")
        document.add_heading(
            spec.text(block.get("text"), f"{label} text", max_chars=512), level=level
        )
        return {"headings": 1}
    if kind == "paragraph":
        document.add_paragraph(spec.text(block.get("text"), f"{label} text"))
        return {"paragraphs": 1}
    if kind in ("bullets", "numbered"):
        style = "List Bullet" if kind == "bullets" else "List Number"
        for item in spec.items(block.get("items"), f"{label} items"):
            document.add_paragraph(spec.text(item, f"{label} item"), style=style)
        return {"lists": 1}
    if kind == "table":
        _table(document, block, label)
        return {"tables": 1}
    if kind == "page_break":
        document.add_page_break()
        return {"page_breaks": 1}
    if kind == "chart":
        chart = spec.chart(block.get("chart"), f"{label} chart")
        document.add_picture(io.BytesIO(chart_render.render(chart)), width=Cm(15.5))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return {"charts": 1}
    raise ValueError(f"{label} has an unknown type: {kind!r}")


def build(content: dict[str, Any], path: Path) -> dict[str, Any]:
    """Write the DOCX to `path` and return what it contains, for the result document."""
    title = spec.text(content.get("title"), "title", max_chars=512)
    document = Document()
    _style_document(document)
    properties = document.core_properties
    properties.title = title
    properties.author = "Delta"
    properties.created = properties.modified = datetime.now(timezone.utc).replace(
        tzinfo=None
    )
    document.add_heading(title, level=0)
    if content.get("subtitle"):
        document.add_paragraph(
            spec.text(content["subtitle"], "subtitle", max_chars=512)
        )
    counts: dict[str, int] = {}
    for index, block in enumerate(spec.blocks(content.get("blocks", []))):
        for key, value in _block(document, block, index).items():
            counts[key] = counts.get(key, 0) + value
    _header_and_footer(
        document,
        spec.text(content.get("header"), "header", is_required=False, max_chars=512),
        spec.text(content.get("footer"), "footer", is_required=False, max_chars=512),
    )
    document.save(str(path))
    return counts
