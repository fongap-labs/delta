"""Capability worker for the Office Worker (ADR-0085).

The Rust host writes one `CapabilityJob` as a JSON line to stdin, then one line of granted secret
values, and reads one `CapabilityResult` as JSON from stdout. The worker writes its candidate file
into the staging directory the host allocated; the host then checks the path, hashes the file and
registers it as an artifact. The worker never writes into the workspace itself.

Run with `python -I worker.py`. The worker adds its own package root to `sys.path`.
"""

from __future__ import annotations

import hashlib
import json
import sys
import time
from pathlib import Path
from typing import Any

_PACK_ROOT = Path(__file__).resolve().parents[2]
if str(_PACK_ROOT) not in sys.path:
    sys.path.insert(0, str(_PACK_ROOT))

from packages.office_worker import docx_build, pptx_build, xlsx_build  # noqa: E402

ABI_VERSION = 2
CAPABILITIES = {
    "office.docx.create": "docx",
    "office.pptx.create": "pptx",
    "office.xlsx.create": "xlsx",
}
BUILDERS = {
    "docx": docx_build.build,
    "pptx": pptx_build.build,
    "xlsx": xlsx_build.build,
}
MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class JobError(Exception):
    """A job the worker refuses, with the machine-readable code the host records."""

    def __init__(self, message: str, code: str) -> None:
        super().__init__(message)
        self.code = code


def _failed(job_id: str, message: str, code: str) -> dict[str, Any]:
    return {
        "abi_version": ABI_VERSION,
        "job_id": job_id,
        "state": "failed",
        "result": None,
        "output": None,
        "artifacts": [],
        "diagnostics": {
            "stderr_lines": [],
            "error_code": code,
            "error_message": message,
        },
        "finished_at": time.time(),
    }


def _completed(
    job_id: str, document: dict[str, Any], artifact: dict[str, Any]
) -> dict[str, Any]:
    return {
        "abi_version": ABI_VERSION,
        "job_id": job_id,
        "state": "completed",
        "result": document,
        "output": None,
        "artifacts": [artifact],
        "diagnostics": None,
        "finished_at": time.time(),
    }


def _verify(kind: str, path: Path, counts: dict[str, int]) -> dict[str, Any]:
    """Reopen the written file and check it holds what the builder reported. Refuse otherwise."""
    try:
        if kind == "docx":
            from docx import Document

            document = Document(str(path))
            headings = sum(
                1
                for paragraph in document.paragraphs
                if docx_build.is_heading(paragraph)
            )
            observed = {"tables": len(document.tables), "headings": headings}
        elif kind == "pptx":
            from pptx import Presentation

            observed = {"slides": len(Presentation(str(path)).slides)}
        else:
            from openpyxl import load_workbook

            observed = {
                "sheets": len(load_workbook(str(path), read_only=True).sheetnames)
            }
    except Exception as exc:  # any unreadable file is a failed verification
        raise JobError(
            f"the written {kind} file cannot be reopened: {exc}", "verify_failed"
        ) from exc
    for key, expected in counts.items():
        if key in observed and observed[key] != expected:
            raise JobError(
                f"the written {kind} file has {observed[key]} {key}, expected {expected}",
                "verify_failed",
            )
    return {"status": "passed", "observed": observed}


def _staging_file(job: dict[str, Any], kind: str) -> Path:
    staging = job.get("artifact_staging_dir")
    if not staging:
        raise JobError("artifact staging is unavailable", "staging")
    directory = Path(str(staging))
    if not directory.is_absolute() or not directory.is_dir():
        raise JobError(
            "artifact staging is not an existing absolute directory", "staging"
        )
    # The extension matters: the readers that verify the file select their format by it.
    return directory / f"candidate.{kind}"


def run_job(job: dict[str, Any]) -> dict[str, Any]:
    """Run one job and return the `CapabilityResult` document as a dictionary."""
    job_id = str(job.get("job_id", ""))
    if job.get("abi_version") != ABI_VERSION:
        return _failed(job_id, "worker ABI version mismatch", "abi_mismatch")
    kind = CAPABILITIES.get(str(job.get("capability_id", "")))
    if kind is None:
        return _failed(
            job_id, "unknown capability for the office worker", "unknown_capability"
        )
    arguments = job.get("arguments") or {}
    relative_path = str(arguments.get("path") or "")
    if not relative_path:
        return _failed(job_id, "path is required", "arguments")
    if not relative_path.endswith(f".{kind}"):
        return _failed(job_id, f"path must end with .{kind}", "arguments")
    try:
        content = arguments.get("content")
        if not isinstance(content, dict):
            raise JobError("content must be a structured spec object", "arguments")
        candidate = _staging_file(job, kind)
        counts = BUILDERS[kind](content, candidate)
        verification = _verify(kind, candidate, counts)
    except JobError as exc:
        return _failed(job_id, str(exc), exc.code)
    except (ValueError, KeyError, TypeError) as exc:
        return _failed(job_id, str(exc), "invalid_spec")
    except OSError as exc:
        return _failed(job_id, f"could not write the candidate: {exc}", "write_failed")
    data = candidate.read_bytes()
    document = {
        "format": kind,
        "path": relative_path,
        "counts": counts,
        "spec_sha256": hashlib.sha256(
            json.dumps(content, sort_keys=True, ensure_ascii=False).encode("utf-8")
        ).hexdigest(),
        "media_type": MEDIA_TYPES[kind],
        "verification": verification,
    }
    artifact = {
        "staging_path": str(candidate),
        "relative_path": relative_path,
        "kind": "file",
        "sha256": hashlib.sha256(data).hexdigest(),
        "size": len(data),
        "incomplete": False,
    }
    return _completed(job_id, document, artifact)


def main() -> None:
    # The host writes the job as the first line and the granted secret values as the second. This
    # worker is granted no secrets, so the second line is read and discarded.
    raw = sys.stdin.readline()
    sys.stdin.read()
    try:
        job = json.loads(raw)
    except json.JSONDecodeError as exc:
        result = _failed("", f"job is not valid JSON: {exc}", "invalid_job")
    else:
        result = run_job(job)
    sys.stdout.write(json.dumps(result, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    main()
