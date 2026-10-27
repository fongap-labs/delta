"""Builds a finished XLSX from a structured spec (ADR-0085).

Each sheet has a styled header row that stays visible while scrolling, an autofilter, column widths
fitted to the content, and numbers stored as numbers. A text cell that begins with "=" is a
formula. Formulas that reach outside the workbook (external references, network or indirect
functions) are refused, because a deliverable must not change when it is opened elsewhere.
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from packages.office_worker import spec

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MIN_WIDTH = 8
MAX_WIDTH = 60
FORMULA_LIMIT = 8_192
FORBIDDEN_IN_FORMULA = re.compile(
    r"\[|\b(?:HYPERLINK|WEBSERVICE|INDIRECT|OFFSET|FILEPATH|CELL|INFO)\s*\(",
    re.IGNORECASE,
)
SHEET_NAME_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(raw: Any, label: str, taken: set[str]) -> str:
    name = spec.text(raw, f"{label} name", max_chars=31)
    if SHEET_NAME_FORBIDDEN.search(name):
        raise ValueError(f"{label} name contains a character a sheet name cannot hold")
    if name.lower() in taken:
        raise ValueError(f"{label} name {name!r} is used by another sheet")
    taken.add(name.lower())
    return name


def _formula(value: str, label: str) -> str:
    if len(value) > FORMULA_LIMIT:
        raise ValueError(f"{label} formula is longer than {FORMULA_LIMIT} characters")
    if FORBIDDEN_IN_FORMULA.search(value):
        raise ValueError(
            f"{label} formula reaches outside the workbook, which is refused"
        )
    return value


def _fill_sheet(
    sheet: Any,
    headers: list[str],
    rows: list[list[Any]],
    formats: dict[str, str],
    label: str,
) -> None:
    sheet.append(headers)
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for row in rows:
        sheet.append(
            [
                _formula(value, f"{label} cell")
                if isinstance(value, str) and value.startswith("=")
                else value
                for value in row
            ]
        )
    for index, name in enumerate(headers, start=1):
        number_format = formats.get(name)
        if number_format:
            for row_index in range(2, len(rows) + 2):
                sheet.cell(row=row_index, column=index).number_format = number_format
    for index, name in enumerate(headers, start=1):
        widest = max(
            [len(name)]
            + [
                len(spec.display(row[index - 1]))
                for row in rows
                if index - 1 < len(row)
            ]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(widest + 2, MIN_WIDTH), MAX_WIDTH
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    )


def build(content: dict[str, Any], path: Path) -> dict[str, Any]:
    """Write the XLSX to `path` and return what it contains, for the result document."""
    title = spec.text(content.get("title"), "title", max_chars=512)
    sheets = spec.items(content.get("sheets"), "sheets", limit=spec.MAX_SHEETS)
    if not sheets:
        raise ValueError("sheets must contain at least one sheet")
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    workbook.properties.title = title
    workbook.properties.creator = "Delta"
    workbook.properties.created = workbook.properties.modified = datetime.now(
        timezone.utc
    ).replace(tzinfo=None)
    taken: set[str] = set()
    cells = 0
    for index, sheet_spec in enumerate(sheets):
        label = f"sheets[{index}]"
        if not isinstance(sheet_spec, dict):
            raise ValueError(f"{label} must be an object")
        name = _sheet_name(sheet_spec.get("name"), label, taken)
        headers, rows = spec.table(
            sheet_spec.get("columns"), sheet_spec.get("rows"), label
        )
        formats = sheet_spec.get("number_formats") or {}
        if not isinstance(formats, dict) or not all(
            isinstance(value, str) for value in formats.values()
        ):
            raise ValueError(
                f"{label} number_formats must map column names to format strings"
            )
        sheet = workbook.create_sheet(title=name)
        _fill_sheet(sheet, headers, rows, formats, label)
        cells += (len(rows) + 1) * len(headers)
    workbook.save(str(path))
    return {"sheets": len(sheets), "cells": cells}
