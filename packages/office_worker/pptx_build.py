"""Builds a finished PPTX from a structured spec (ADR-0085).

Slides are 16:9. A title slide, a bullet slide and a table slide cover the common cases; every
slide may carry speaker notes. Text sizes are fixed per role, so a deck reads the same whatever
its content is.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from packages.office_worker import spec

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)
# The default template is 4:3 (10 inches wide). Its placeholders are scaled to the 16:9 width.
TEMPLATE_WIDTH = Inches(10)
TITLE_LAYOUT = 0
CONTENT_LAYOUT = 1
TITLE_ONLY_LAYOUT = 5
LAYOUTS = ("title", "bullets", "table")
BULLET_SIZE = Pt(20)
TABLE_SIZE = Pt(14)
TITLE_SIZE = Pt(32)
HEADING_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT_HEIGHT = Inches(0.12)


def _fit_to_width(slide: Any) -> None:
    factor = SLIDE_WIDTH / TEMPLATE_WIDTH
    for shape in slide.placeholders:
        # Set all four geometry values: an inherited placeholder that gets only a left and width
        # would otherwise be written with a top and height of zero.
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        shape.left, shape.top = int(left * factor), top
        shape.width, shape.height = int(width * factor), height


def _decorate(slide: Any) -> None:
    """One house style on every slide: a navy title and a thin accent bar along the top."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, ACCENT_HEIGHT)
    bar.fill.solid()
    bar.fill.fore_color.rgb = HEADING_COLOR
    bar.line.fill.background()
    title = slide.shapes.title
    if title is not None and title.has_text_frame:
        for paragraph in title.text_frame.paragraphs:
            paragraph.font.size = TITLE_SIZE
            paragraph.font.bold = True
            paragraph.font.color.rgb = HEADING_COLOR


def _title_slide(presentation: Any, slide_spec: dict[str, Any], label: str) -> Any:
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_LAYOUT])
    slide.shapes.title.text = spec.text(
        slide_spec.get("title"), f"{label} title", max_chars=512
    )
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = spec.text(
            slide_spec.get("subtitle"),
            f"{label} subtitle",
            required=False,
            max_chars=512,
        )
    return slide


def _bullet_slide(presentation: Any, slide_spec: dict[str, Any], label: str) -> Any:
    slide = presentation.slides.add_slide(presentation.slide_layouts[CONTENT_LAYOUT])
    slide.shapes.title.text = spec.text(
        slide_spec.get("title"), f"{label} title", max_chars=512
    )
    frame = slide.placeholders[1].text_frame
    frame.clear()
    for index, item in enumerate(
        spec.items(slide_spec.get("items"), f"{label} items", limit=200)
    ):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = spec.text(item, f"{label} item")
        paragraph.font.size = BULLET_SIZE
        paragraph.level = 0
    return slide


def _table_slide(presentation: Any, slide_spec: dict[str, Any], label: str) -> Any:
    headers, rows = spec.table(slide_spec.get("columns"), slide_spec.get("rows"), label)
    slide = presentation.slides.add_slide(presentation.slide_layouts[TITLE_ONLY_LAYOUT])
    slide.shapes.title.text = spec.text(
        slide_spec.get("title"), f"{label} title", max_chars=512
    )
    left, top = Inches(0.6), Inches(1.8)
    width = SLIDE_WIDTH - Inches(1.2)
    height = Emu(min(len(rows) + 1, 25) * Inches(0.5))
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), left, top, width, height
    )
    grid = shape.table
    for column, name in enumerate(headers):
        grid.cell(0, column).text = name
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            cell = grid.cell(row_index, column)
            cell.text = spec.display(value)
            if spec.is_number(value):
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    for row_index, row in enumerate(grid.rows):
        for cell in row.cells:
            if row_index == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = HEADING_COLOR
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = TABLE_SIZE
                paragraph.font.bold = row_index == 0
                if row_index == 0:
                    paragraph.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return slide


def build(content: dict[str, Any], path: Path) -> dict[str, Any]:
    """Write the PPTX to `path` and return what it contains, for the result document."""
    title = spec.text(content.get("title"), "title", max_chars=512)
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    properties = presentation.core_properties
    properties.title = title
    properties.author = "Delta"
    properties.created = properties.modified = datetime.now(timezone.utc).replace(
        tzinfo=None
    )
    counts: dict[str, int] = {"slides": 0, "notes": 0}
    slides = spec.items(content.get("slides"), "slides", limit=spec.MAX_SLIDES)
    if not slides:
        raise ValueError("slides must contain at least one slide")
    for index, slide_spec in enumerate(slides):
        label = f"slides[{index}]"
        if not isinstance(slide_spec, dict):
            raise ValueError(f"{label} must be an object")
        layout = slide_spec.get("layout")
        if layout == "title":
            slide = _title_slide(presentation, slide_spec, label)
        elif layout == "bullets":
            slide = _bullet_slide(presentation, slide_spec, label)
        elif layout == "table":
            slide = _table_slide(presentation, slide_spec, label)
        else:
            raise ValueError(f"{label} layout must be one of {', '.join(LAYOUTS)}")
        _fit_to_width(slide)
        _decorate(slide)
        counts["slides"] += 1
        if slide_spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec.text(
                slide_spec["notes"], f"{label} notes"
            )
            counts["notes"] += 1
    presentation.save(str(path))
    return counts
